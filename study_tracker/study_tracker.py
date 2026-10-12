from __future__ import annotations

import os
import sys
import tempfile
import tkinter as tk
from dataclasses import dataclass, replace
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook


DONE_MARK = "✔"
TOPIC_HEADERS = ("Done", "#", "Topic", "Subtopic", "What to cover", "Notes")


@dataclass(frozen=True)
class Topic:
    area: str
    row: int
    number: int
    title: str
    subtopic: str
    coverage: str
    notes: str
    done: bool


def get_default_workbook_path() -> Path:
    return Path.home() / "OneDrive" / "Desktop" / "Study_Tracker.xlsx"


def get_icon_path() -> Path:
    bundle_directory = Path(getattr(sys, "_MEIPASS", Path(__file__).parent))
    return bundle_directory / "study_tracker.ico"


def load_topics(workbook_path: Path) -> list[Topic]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=False)
    topics: list[Topic] = []
    try:
        for worksheet in workbook.worksheets:
            headers = tuple(
                worksheet.cell(row=1, column=column).value
                for column in range(1, len(TOPIC_HEADERS) + 1)
            )
            if headers != TOPIC_HEADERS:
                continue

            for row_number in range(2, worksheet.max_row + 1):
                values = tuple(
                    worksheet.cell(row=row_number, column=column).value
                    for column in range(1, len(TOPIC_HEADERS) + 1)
                )
                if not any(value is not None for value in values):
                    continue

                topic_number = values[1]
                topics.append(
                    Topic(
                        area=worksheet.title,
                        row=row_number,
                        number=(
                            int(topic_number)
                            if isinstance(topic_number, (int, float))
                            else row_number - 1
                        ),
                        title=str(values[2] or "").strip(),
                        subtopic=str(values[3] or "").strip(),
                        coverage=str(values[4] or "").strip(),
                        notes=str(values[5] or "").strip(),
                        done=values[0] == DONE_MARK,
                    )
                )
    finally:
        workbook.close()

    if not topics:
        raise ValueError(
            "No topic sheets were found. Each topic sheet must have the "
            "columns Done, #, Topic, Subtopic, What to cover, and Notes."
        )
    return topics


def save_topic_completion(
    workbook_path: Path,
    topic: Topic,
    completed: bool,
) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    temporary_path: Path | None = None
    try:
        worksheet = workbook[topic.area]
        worksheet.cell(row=topic.row, column=1).value = DONE_MARK if completed else None

        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_path.stem}_",
            suffix=workbook_path.suffix,
            dir=workbook_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


class StudyTrackerApp:
    def __init__(self, root: tk.Tk, workbook_path: Path, topics: list[Topic]) -> None:
        self.root = root
        self.workbook_path = workbook_path
        self.topics = topics
        self.topic_by_key = {(topic.area, topic.row): topic for topic in topics}
        self.search_value = tk.StringVar()
        self.area_value = tk.StringVar(value="All areas")
        self.workbook_label = tk.StringVar(value=str(workbook_path))
        self.status_value = tk.StringVar(value="Changes are saved to the workbook.")
        self.progress_values: dict[str, tk.StringVar] = {}

        root.title("Study Tracker")
        root.geometry("1120x740")
        root.minsize(900, 580)
        root.configure(bg="#f4f6fa")

        self._build_interface()
        self.refresh_topics()

    def _build_interface(self) -> None:
        style = ttk.Style(self.root)
        style.configure("Treeview", rowheight=30, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))

        container = ttk.Frame(self.root, padding=20)
        container.pack(fill="both", expand=True)
        container.columnconfigure(0, weight=1)
        container.rowconfigure(3, weight=1)

        heading = ttk.Frame(container)
        heading.grid(row=0, column=0, sticky="ew")
        heading.columnconfigure(0, weight=1)
        ttk.Label(
            heading,
            text="Study Tracker",
            font=("Segoe UI", 22, "bold"),
        ).grid(row=0, column=0, sticky="w")
        ttk.Button(
            heading,
            text="Choose workbook...",
            command=self.choose_workbook,
        ).grid(row=0, column=1, sticky="e")
        ttk.Label(
            heading,
            textvariable=self.workbook_label,
            foreground="#687386",
        ).grid(row=1, column=0, columnspan=2, sticky="w", pady=(2, 0))

        self.progress_frame = ttk.Frame(container)
        self.progress_frame.grid(row=1, column=0, sticky="ew", pady=(18, 14))
        self.progress_frame.columnconfigure(0, weight=1)

        filters = ttk.Frame(container)
        filters.grid(row=2, column=0, sticky="ew", pady=(0, 12))
        filters.columnconfigure(3, weight=1)
        ttk.Label(filters, text="Area").grid(row=0, column=0, padx=(0, 8))
        areas = ["All areas", *dict.fromkeys(topic.area for topic in self.topics)]
        self.area_combo = ttk.Combobox(
            filters,
            textvariable=self.area_value,
            values=areas,
            state="readonly",
            width=18,
        )
        self.area_combo.grid(row=0, column=1, padx=(0, 18))
        self.area_combo.bind("<<ComboboxSelected>>", lambda _event: self.refresh_topics())
        ttk.Label(filters, text="Search").grid(row=0, column=2, padx=(0, 8))
        search_entry = ttk.Entry(filters, textvariable=self.search_value)
        search_entry.grid(row=0, column=3, sticky="ew")
        self.search_value.trace_add("write", lambda *_args: self.refresh_topics())
        self.root.bind("<Control-f>", lambda _event: search_entry.focus_set())

        content = ttk.Frame(container)
        content.grid(row=3, column=0, sticky="nsew")
        content.columnconfigure(0, weight=3)
        content.columnconfigure(1, weight=2)
        content.rowconfigure(0, weight=1)

        list_frame = ttk.Frame(content)
        list_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        self.topic_tree = ttk.Treeview(
            list_frame,
            columns=("status", "area", "topic", "subtopic"),
            show="headings",
            selectmode="browse",
        )
        for column, title, width in (
            ("status", "Status", 82),
            ("area", "Area", 85),
            ("topic", "Topic", 210),
            ("subtopic", "Subtopic", 150),
        ):
            self.topic_tree.heading(column, text=title)
            self.topic_tree.column(column, width=width, minwidth=60, stretch=column == "topic")
        self.topic_tree.tag_configure("done", foreground="#18804b")
        self.topic_tree.tag_configure("todo", foreground="#30394a")
        self.topic_tree.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(
            list_frame,
            orient="vertical",
            command=self.topic_tree.yview,
        )
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.topic_tree.configure(yscrollcommand=scrollbar.set)
        self.topic_tree.bind("<<TreeviewSelect>>", self.show_selected_topic)
        self.topic_tree.bind("<Double-1>", self.toggle_selected)

        detail_frame = ttk.LabelFrame(content, text="Topic details", padding=16)
        detail_frame.grid(row=0, column=1, sticky="nsew")
        detail_frame.columnconfigure(0, weight=1)
        detail_frame.rowconfigure(3, weight=1)
        self.detail_title = ttk.Label(
            detail_frame,
            text="Select a topic",
            font=("Segoe UI", 15, "bold"),
            wraplength=340,
        )
        self.detail_title.grid(row=0, column=0, sticky="w")
        self.detail_subtitle = ttk.Label(
            detail_frame,
            text="",
            foreground="#687386",
            wraplength=340,
        )
        self.detail_subtitle.grid(row=1, column=0, sticky="w", pady=(4, 12))
        ttk.Label(detail_frame, text="What to cover", font=("Segoe UI", 10, "bold")).grid(
            row=2, column=0, sticky="nw"
        )
        self.coverage_text = self._create_readonly_text(detail_frame)
        self.coverage_text.grid(row=3, column=0, sticky="nsew", pady=(5, 14))
        ttk.Label(detail_frame, text="Notes", font=("Segoe UI", 10, "bold")).grid(
            row=4, column=0, sticky="w"
        )
        self.notes_text = self._create_readonly_text(detail_frame, height=5)
        self.notes_text.grid(row=5, column=0, sticky="ew", pady=(5, 14))
        self.toggle_button = ttk.Button(
            detail_frame,
            text="Mark as done",
            command=self.toggle_selected,
            state="disabled",
        )
        self.toggle_button.grid(row=6, column=0, sticky="ew")

        footer = ttk.Frame(container)
        footer.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        ttk.Label(footer, textvariable=self.status_value, foreground="#687386").pack(
            side="left"
        )
        ttk.Label(footer, text="Tip: double-click a topic to toggle completion").pack(
            side="right"
        )

    @staticmethod
    def _create_readonly_text(parent: ttk.Frame, height: int = 8) -> tk.Text:
        text = tk.Text(
            parent,
            height=height,
            wrap="word",
            font=("Segoe UI", 10),
            relief="solid",
            borderwidth=1,
            padx=9,
            pady=8,
            background="#fbfcfe",
        )
        text.configure(state="disabled")
        return text

    @staticmethod
    def _set_text(widget: tk.Text, value: str) -> None:
        widget.configure(state="normal")
        widget.delete("1.0", tk.END)
        widget.insert("1.0", value)
        widget.configure(state="disabled")

    def _update_progress(self) -> None:
        for child in self.progress_frame.winfo_children():
            child.destroy()

        areas = list(dict.fromkeys(topic.area for topic in self.topics))
        for column, area in enumerate(["Overall", *areas]):
            area_topics = self.topics if area == "Overall" else [
                topic for topic in self.topics if topic.area == area
            ]
            completed = sum(topic.done for topic in area_topics)
            total = len(area_topics)
            progress = f"{completed}/{total} done  ·  {completed / total:.0%}" if total else "0/0 done"
            card = ttk.LabelFrame(self.progress_frame, text=area, padding=(12, 8))
            card.grid(row=0, column=column, sticky="ew", padx=(0, 8))
            card.columnconfigure(0, weight=1)
            variable = self.progress_values.setdefault(area, tk.StringVar())
            variable.set(progress)
            ttk.Label(card, textvariable=variable, font=("Segoe UI", 10, "bold")).grid(
                row=0, column=0, sticky="w"
            )
            self.progress_frame.columnconfigure(column, weight=1)

    def refresh_topics(self) -> None:
        selected = self.topic_tree.selection() if hasattr(self, "topic_tree") else ()
        selected_key = selected[0] if selected else None
        for item in self.topic_tree.get_children():
            self.topic_tree.delete(item)

        search = self.search_value.get().strip().casefold()
        selected_area = self.area_value.get()
        visible_topics = []
        for topic in self.topics:
            if selected_area != "All areas" and topic.area != selected_area:
                continue
            searchable = " ".join(
                (topic.area, topic.title, topic.subtopic, topic.coverage, topic.notes)
            ).casefold()
            if search and search not in searchable:
                continue
            visible_topics.append(topic)
            key = f"{topic.area}|{topic.row}"
            self.topic_tree.insert(
                "",
                "end",
                iid=key,
                values=(
                    "✔ Done" if topic.done else "To do",
                    topic.area,
                    topic.title,
                    topic.subtopic,
                ),
                tags=("done" if topic.done else "todo",),
            )

        self._update_progress()
        if selected_key and self.topic_tree.exists(selected_key):
            self.topic_tree.selection_set(selected_key)
            self.topic_tree.focus(selected_key)
        elif visible_topics:
            first_key = f"{visible_topics[0].area}|{visible_topics[0].row}"
            self.topic_tree.selection_set(first_key)
            self.topic_tree.focus(first_key)
        else:
            self._clear_details()

    def _clear_details(self) -> None:
        self.detail_title.configure(text="No topics found")
        self.detail_subtitle.configure(text="")
        self._set_text(self.coverage_text, "")
        self._set_text(self.notes_text, "")
        self.toggle_button.configure(text="Mark as done", state="disabled")

    def show_selected_topic(self, _event: tk.Event[tk.Misc] | None = None) -> None:
        selection = self.topic_tree.selection()
        if not selection:
            self._clear_details()
            return

        area, row_text = selection[0].rsplit("|", 1)
        topic = self.topic_by_key[(area, int(row_text))]
        self.detail_title.configure(text=topic.title)
        subtopic = topic.subtopic if topic.subtopic and topic.subtopic != "—" else ""
        self.detail_subtitle.configure(
            text=f"{topic.area}  ·  Topic {topic.number}"
            + (f"  ·  {subtopic}" if subtopic else "")
        )
        self._set_text(self.coverage_text, topic.coverage or "No coverage details.")
        self._set_text(self.notes_text, topic.notes or "No notes.")
        self.toggle_button.configure(
            text="Mark as to do" if topic.done else "Mark as done",
            state="normal",
        )

    def toggle_selected(self, _event: tk.Event[tk.Misc] | None = None) -> None:
        selection = self.topic_tree.selection()
        if not selection:
            return

        area, row_text = selection[0].rsplit("|", 1)
        key = (area, int(row_text))
        topic = self.topic_by_key[key]
        completed = not topic.done
        try:
            save_topic_completion(self.workbook_path, topic, completed)
        except (OSError, KeyError, ValueError) as error:
            messagebox.showerror(
                "Could not save progress",
                f"The workbook was not updated:\n{error}",
                parent=self.root,
            )
            return

        updated_topic = replace(topic, done=completed)
        self.topic_by_key[key] = updated_topic
        self.topics = [
            updated_topic if item.area == area and item.row == key[1] else item
            for item in self.topics
        ]
        self.status_value.set("Progress saved to the workbook.")
        self.refresh_topics()
        self.topic_tree.selection_set(selection[0])
        self.topic_tree.focus(selection[0])
        self.show_selected_topic()

    def choose_workbook(self) -> None:
        selected_path = filedialog.askopenfilename(
            parent=self.root,
            title="Choose a study tracker workbook",
            filetypes=(("Excel workbook", "*.xlsx"),),
            initialdir=str(self.workbook_path.parent),
        )
        if not selected_path:
            return

        workbook_path = Path(selected_path)
        try:
            topics = load_topics(workbook_path)
        except (OSError, ValueError) as error:
            messagebox.showerror(
                "Could not open workbook",
                str(error),
                parent=self.root,
            )
            return

        self.workbook_path = workbook_path
        self.topics = topics
        self.topic_by_key = {(topic.area, topic.row): topic for topic in topics}
        self.workbook_label.set(str(workbook_path))
        self.area_value.set("All areas")
        self.area_combo.configure(
            values=["All areas", *dict.fromkeys(topic.area for topic in topics)]
        )
        self.status_value.set("Workbook loaded.")
        self.refresh_topics()


def main() -> None:
    root = tk.Tk()
    root.iconbitmap(str(get_icon_path()))
    workbook_path = get_default_workbook_path()
    if not workbook_path.is_file():
        selected_path = filedialog.askopenfilename(
            parent=root,
            title="Choose your Study_Tracker.xlsx workbook",
            filetypes=(("Excel workbook", "*.xlsx"),),
            initialdir=str(Path.home() / "OneDrive" / "Desktop"),
        )
        if not selected_path:
            root.destroy()
            return
        workbook_path = Path(selected_path)

    try:
        topics = load_topics(workbook_path)
    except (OSError, ValueError) as error:
        messagebox.showerror("Study Tracker", str(error), parent=root)
        root.destroy()
        return

    StudyTrackerApp(root, workbook_path, topics)
    root.mainloop()


if __name__ == "__main__":
    main()
