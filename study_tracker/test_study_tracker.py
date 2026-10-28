import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from study_tracker import (
    Topic,
    add_topic,
    delete_topic,
    get_workbook_error_message,
    load_topic_areas,
    load_topics,
    save_topic_completion,
)


class StudyTrackerWorkbookTests(unittest.TestCase):
    def create_workbook(self, path: Path) -> None:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary["A1"] = "Study Progress Tracker"
        for column, value in enumerate(
            ("Area", "Total", "Completed", "Remaining", "Progress"),
            start=1,
        ):
            summary.cell(row=4, column=column, value=value)
        for column, value in enumerate(
            ("Python", "=COUNTA(Python!C2:C2)", '=COUNTIF(Python!A2:A2,"✔")', "=B5-C5", "=IF(B5=0,0,C5/B5)"),
            start=1,
        ):
            summary.cell(row=5, column=column, value=value)
        for column, value in enumerate(
            ("Overall", "=SUM(B5)", "=SUM(C5)", "=SUM(D5)", "=IF(B6=0,0,C6/B6)"),
            start=1,
        ):
            summary.cell(row=6, column=column, value=value)
        sheet = workbook.create_sheet("Python")
        sheet.append(("Done", "#", "Topic", "Subtopic", "What to cover", "Notes"))
        sheet.append((None, 1, "Functions", "Arguments", "Parameters and values", "Review"))
        workbook.save(path)
        workbook.close()

    def test_load_topics_reads_fields_and_completion_status(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)
            workbook = load_workbook(path)
            workbook["Python"]["A2"] = "✔"
            workbook.save(path)
            workbook.close()

            topics = load_topics(path)

        self.assertEqual(len(topics), 1)
        self.assertEqual(
            topics[0],
            Topic(
                area="Python",
                row=2,
                number=1,
                title="Functions",
                subtopic="Arguments",
                coverage="Parameters and values",
                notes="Review",
                done=True,
            ),
        )

    def test_completion_is_saved_without_overwriting_summary_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)
            workbook = load_workbook(path)
            workbook["Summary"]["A2"] = "=COUNTA(Python!C2:C2)"
            workbook.save(path)
            workbook.close()

            topic = load_topics(path)[0]
            save_topic_completion(path, topic, completed=True)

            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook["Python"]["A2"].value, "✔")
            self.assertEqual(workbook["Summary"]["A2"].value, "=COUNTA(Python!C2:C2)")
            workbook.close()

            self.assertTrue(load_topics(path)[0].done)
            save_topic_completion(path, topic, completed=False)
            self.assertFalse(load_topics(path)[0].done)

    def test_missing_workbook_is_reported(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            missing_path = Path(temporary_directory) / "missing.xlsx"

            with self.assertRaisesRegex(FileNotFoundError, "Workbook not found"):
                load_topics(missing_path)

    def test_permission_error_explains_recovery_without_discarding_entry(self) -> None:
        error = PermissionError(13, "Permission denied", r"C:\Study_Tracker.xlsx")

        message = get_workbook_error_message("add the topic", error)

        self.assertIn("Windows could not add the topic", message)
        self.assertIn(r"C:\Study_Tracker.xlsx", message)
        self.assertIn("Your changes are still in the app", message)
        self.assertIn("close the workbook in Excel", message)
        self.assertIn("OneDrive", message)
        self.assertIn("Choose workbook...", message)

    def test_add_topic_appends_fields_and_updates_summary_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)

            add_topic(path, "Python", "Decorators", "Wrappers", "Wrapper functions", "Practice")

            topics = load_topics(path)
            workbook = load_workbook(path, data_only=False)
            summary = workbook["Summary"]
            self.assertEqual(len(topics), 2)
            self.assertEqual(
                topics[-1],
                Topic(
                    area="Python",
                    row=3,
                    number=2,
                    title="Decorators",
                    subtopic="Wrappers",
                    coverage="Wrapper functions",
                    notes="Practice",
                    done=False,
                ),
            )
            self.assertEqual(summary["B5"].value, "=COUNTA('Python'!C2:C3)")
            self.assertEqual(summary["C5"].value, '=COUNTIF(\'Python\'!A2:A3,"✔")')
            self.assertEqual(summary["B6"].value, "=SUM(B5)")
            workbook.close()

    def test_add_topic_requires_a_title(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)

            with self.assertRaisesRegex(ValueError, "Topic is required"):
                add_topic(path, "Python", "  ", "", "", "")

    def test_delete_topic_removes_row_renumbers_and_updates_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)
            add_topic(path, "Python", "Decorators", "Wrappers", "Wrapper functions", "Practice")
            topic_to_delete = load_topics(path)[0]

            delete_topic(path, topic_to_delete)

            topics = load_topics(path)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(len(topics), 1)
            self.assertEqual(topics[0].title, "Decorators")
            self.assertEqual(topics[0].number, 1)
            self.assertEqual(workbook["Summary"]["B5"].value, "=COUNTA('Python'!C2:C2)")
            self.assertEqual(
                workbook["Summary"]["C5"].value,
                '=COUNTIF(\'Python\'!A2:A2,"✔")',
            )
            workbook.close()

    def test_deleting_last_topic_keeps_area_available_for_new_topics(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "tracker.xlsx"
            self.create_workbook(path)
            delete_topic(path, load_topics(path)[0])

            self.assertEqual(load_topics(path), [])
            self.assertEqual(load_topic_areas(path), ["Python"])
            add_topic(path, "Python", "Generators", "", "", "")
            self.assertEqual([topic.title for topic in load_topics(path)], ["Generators"])


if __name__ == "__main__":
    unittest.main()
